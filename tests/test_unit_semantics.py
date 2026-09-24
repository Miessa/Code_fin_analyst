# -*- coding: utf-8 -*-
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

from arsel_core.collecter_libelles import collecter
from arsel_core.unit_semantics import (
    choose_unit_for_metric,
    compatibility,
    decompose_unit,
)


class TestUnitSemantics(unittest.TestCase):
    def test_decompose_devise_echelle_et_periodicite(self):
        unit = decompose_unit("EUR'000s/an")
        self.assertEqual(unit["currency"], "EUR")
        self.assertEqual(unit["scale"], 1000)
        self.assertEqual(unit["periodicity"], "annual")
        self.assertEqual(unit["family"], "currency_per_year")

    def test_distingue_puissance_et_energie(self):
        self.assertEqual(decompose_unit("MW")["family"], "power")
        self.assertEqual(decompose_unit("MWh")["family"], "energy")

    def test_tarif_compose_devise_puissance_et_mois(self):
        unit = decompose_unit("Base (EUR/kW/month)")
        self.assertEqual(unit["family"], "tariff_rate")
        self.assertEqual(unit["currency"], "EUR")
        self.assertEqual(unit["denominator"], "kW")
        self.assertEqual(unit["periodicity"], "monthly")
        self.assertEqual(unit["unit_expression"], "EUR /kW /month")

    def test_dscr_refuse_un_pourcentage(self):
        concept = {"famille_unite": "multiple_ratio"}
        self.assertFalse(compatibility(decompose_unit("% p.a."), concept)["compatible"])
        self.assertTrue(compatibility(decompose_unit("x"), concept)["compatible"])

    def test_compatibilite_depend_de_la_metrique(self):
        unit = decompose_unit("MW")
        self.assertTrue(compatibility(unit, {"famille_unite": "power"})["compatible"])
        self.assertFalse(compatibility(unit, {"famille_unite": "duration_years"})["compatible"])

    def test_mauvaise_unite_ne_disqualifie_pas_project_cost(self):
        candidate = {
            "libelle": "Project Cost",
            "valeur": 2472419,
            "score_retrieval_final": 0.95,
            "unit_candidates": [{
                "source_text": "years",
                "source_cell": "InpC!G1107",
                "spatial_relation": "same_row_right",
                "attachment_confidence": 0.98,
                "semantics": decompose_unit("years"),
            }],
        }
        result = choose_unit_for_metric(
            candidate, {"cle": "investissement_total", "famille_unite": "currency_amount"}
        )
        self.assertEqual(result["score_retrieval_final"], 0.95)
        self.assertEqual(result["unit_status"], "conflict_metric_kept")
        self.assertFalse(result["unit_compatibility"]["compatible"])
        self.assertIsNone(result["unite_detectee"])

    def test_unite_compatible_est_preferee_a_une_unite_incompatible_plus_proche(self):
        candidate = {
            "unit_candidates": [
                {
                    "source_text": "years", "source_cell": "S!C2",
                    "spatial_relation": "same_row_right", "attachment_confidence": 0.98,
                    "semantics": decompose_unit("years"),
                },
                {
                    "source_text": "EUR'000s", "source_cell": "S!B1",
                    "spatial_relation": "upper_header", "attachment_confidence": 0.84,
                    "semantics": decompose_unit("EUR'000s"),
                },
            ]
        }
        result = choose_unit_for_metric(candidate, {"famille_unite": "currency_amount"})
        self.assertEqual(result["unit_semantics"]["currency"], "EUR")
        self.assertEqual(result["unit_semantics"]["scale"], 1000)
        self.assertEqual(result["unit_status"], "confirmed")


class TestSpatialUnitCollection(unittest.TestCase):
    def _collect(self, workbook):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "units.xlsx"
            workbook.save(path)
            return collecter(path, feuilles=[workbook.active.title])

    def test_detecte_unite_au_dessus(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "Summary"
        ws["B1"] = "EUR'000s"
        ws["A2"] = "Project Cost"
        ws["B2"] = 2472419
        candidate = self._collect(wb)[0]
        relations = {u["spatial_relation"] for u in candidate["unit_candidates"]}
        self.assertIn("upper_header", relations)

    def test_detecte_unite_a_droite(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "Summary"
        ws["A1"] = "Project Cost"
        ws["B1"] = 2472419
        ws["C1"] = "EUR'000s"
        candidate = self._collect(wb)[0]
        self.assertTrue(any(u["spatial_relation"] == "same_row_right" for u in candidate["unit_candidates"]))

    def test_detecte_unite_a_gauche_de_la_valeur(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "Summary"
        ws["A1"] = "Fixed O&M costs"
        ws["B1"] = "XAF'000s"
        ws["C1"] = 53045
        candidate = self._collect(wb)[0]
        self.assertTrue(any(u["spatial_relation"] == "same_row_left" for u in candidate["unit_candidates"]))

    def test_collecte_par_defaut_une_feuille_au_nom_inconnu(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "Financial Model 2026"
        ws["A1"] = "Total project investment"
        ws["B1"] = 1250000
        candidates = self._collect(wb)
        self.assertEqual(len(candidates), 1)
        self.assertEqual(candidates[0]["feuille"], "Financial Model 2026")


if __name__ == "__main__":
    unittest.main()
