# -*- coding: utf-8 -*-
"""Génère la note de cadrage de l'interface utilisateur ARSEL."""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "outputs" / "Cadrage_interface_utilisateur_ARSEL.docx"
NAVY = "17365D"
BLUE = "D9EAF7"
GREEN = "E2F0D9"
AMBER = "FFF2CC"
GREY = "F2F2F2"


def shade(cell, color):
    props = cell._tc.get_or_add_tcPr()
    element = props.find(qn("w:shd"))
    if element is None:
        element = OxmlElement("w:shd")
        props.append(element)
    element.set(qn("w:fill"), color)


def set_cell(cell, value, bold=False, color=None):
    cell.text = ""
    run = cell.paragraphs[0].add_run(str(value))
    run.font.name = "Aptos"
    run.font.size = Pt(9)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def configure(document):
    section = document.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08
    for name, size in (("Title", 24), ("Heading 1", 17), ("Heading 2", 13)):
        style = document.styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(NAVY)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("Projet ARSEL  |  Cadrage de l’interface utilisateur")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(100, 100, 100)


def callout(document, title, content, color=BLUE):
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    shade(cell, color)
    paragraph = cell.paragraphs[0]
    title_run = paragraph.add_run(title + "\n")
    title_run.bold = True
    title_run.font.color.rgb = RGBColor.from_string(NAVY)
    title_run.font.size = Pt(11)
    paragraph.add_run(content)
    document.add_paragraph()


def bullets(document, values):
    for value in values:
        document.add_paragraph(value, style="List Bullet")


def architecture_table(document):
    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for index, title in enumerate(("Couche", "Rôle", "Solution envisageable")):
        shade(table.rows[0].cells[index], NAVY)
        set_cell(table.rows[0].cells[index], title, True, "FFFFFF")
    rows = (
        ("Interface web", "Accès par navigateur, dépôt des fichiers, validation et consultation", "React ou interface web équivalente"),
        ("API applicative", "Authentification, orchestration des traitements et contrôle des droits", "FastAPI / Python"),
        ("Traitements", "Extraction, règles métier, benchmarks et production des rapports", "Moteur ARSEL existant et tâches en arrière-plan"),
        ("Données opérationnelles", "Utilisateurs, projets, décisions et historique", "PostgreSQL"),
        ("Données analytiques", "Projets comparables et statistiques sectorielles", "DuckDB côté serveur"),
        ("Stockage documentaire", "Classeurs, rapports, journaux et snapshots", "Stockage centralisé chiffré"),
    )
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for col_index, value in enumerate(row):
            set_cell(cells[col_index], value, bold=col_index == 0)
            if row_index % 2:
                shade(cells[col_index], GREY)


def build():
    document = Document()
    configure(document)
    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Cadrage de l’interface utilisateur ARSEL")
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Principes de conception humain–machine et humain–IA")
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor.from_string(NAVY)
    document.add_paragraph()
    callout(
        document,
        "Orientation recommandée",
        "Le besoin d’un accès depuis n’importe quel lieu conduit à privilégier une application web sécurisée. DuckDB peut rester utilisé pour les traitements analytiques, mais la base, les fichiers et le moteur doivent être hébergés sur une infrastructure centrale plutôt que sur le poste d’un analyste.",
        GREEN,
    )

    document.add_heading("1. Principe directeur", level=1)
    document.add_paragraph(
        "L’interface doit reproduire le raisonnement métier de l’analyste plutôt que l’organisation interne du code. Elle doit rendre les propositions du système compréhensibles, permettre la correction à tout moment et conserver la responsabilité de la décision finale du côté humain."
    )

    document.add_heading("2. Accès à distance et architecture", level=1)
    document.add_paragraph(
        "Une application web permettrait au superviseur et aux analystes de se connecter depuis un navigateur au bureau ou à domicile. Les traitements continueraient sur le serveur même si l’utilisateur ferme la page. Pour une démonstration rapide, Streamlit reste envisageable; pour une solution institutionnelle durable, une API Python et une interface web séparée offriraient davantage de contrôle, de sécurité et d’évolutivité."
    )
    architecture_table(document)
    document.add_paragraph()
    callout(
        document,
        "Point d’attention",
        "Un accès depuis n’importe où ne signifie pas exposer directement le fichier DuckDB sur Internet. Les utilisateurs interrogent l’application; seule l’application accède aux données selon leurs droits.",
        AMBER,
    )

    document.add_heading("3. Parcours utilisateur", level=1)
    document.add_paragraph("Le parcours proposé suit les étapes réelles de l’analyse :")
    steps = (
        "Créer ou ouvrir un projet.",
        "Déposer le modèle financier Excel.",
        "Renseigner et vérifier le contexte du projet.",
        "Lancer l’extraction et suivre son avancement.",
        "Examiner les métriques et leurs candidats.",
        "Valider, corriger ou déclarer une métrique indisponible.",
        "Lancer les calculs et la comparaison avec les benchmarks.",
        "Examiner les résultats et leurs sources.",
        "Générer puis télécharger le rapport.",
        "Retrouver les versions et décisions antérieures.",
    )
    for index, step in enumerate(steps, 1):
        document.add_paragraph(f"{index}. {step}")

    document.add_heading("4. Écran central : comparaison des candidats", level=1)
    document.add_paragraph(
        "La validation des métriques sera l’écran le plus important. Tous les candidats pertinents devraient être visibles simultanément sous forme de tableau ou de cartes. La proposition principale doit être mise en évidence sans masquer les alternatives."
    )
    table = document.add_table(rows=1, cols=6)
    table.style = "Table Grid"
    headers = ("Candidat", "Valeur", "Unité", "Cellule", "Confiance", "Justification")
    for index, value in enumerate(headers):
        shade(table.rows[0].cells[index], NAVY)
        set_cell(table.rows[0].cells[index], value, True, "FFFFFF")
    examples = (
        ("Project Cost", "2 472 419", "milliers EUR", "InpC!F1107", "Élevée", "Total du projet détecté"),
        ("Total financing cost", "530 710", "milliers EUR", "Summary!M56", "Faible", "Périmètre incomplet"),
    )
    for row in examples:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            set_cell(cells[index], value, bold=index == 0)
    document.add_paragraph()
    bullets(document, (
        "Afficher un aperçu de la feuille autour de la cellule sélectionnée.",
        "Présenter la formule Excel et les cellules dont elle dépend.",
        "Permettre le tri selon le score, la feuille, la valeur ou l’unité.",
        "Expliquer pourquoi un candidat est favorisé ou pénalisé.",
        "Permettre une correction simultanée de la valeur et de l’unité.",
        "Permettre de déclarer la métrique indisponible sans forcer un choix.",
    ))

    document.add_heading("5. Principes de conception humain–IA", level=1)
    document.add_heading("5.1 Rendre l’incertitude visible", level=2)
    document.add_paragraph(
        "L’interface ne doit pas présenter un score technique comme une certitude. Des états compréhensibles sont préférables : proposition fiable, proposition à vérifier, ambiguïté entre plusieurs candidats, unité incertaine, transformation par une règle métier ou absence de candidat fiable."
    )
    document.add_heading("5.2 Fournir une explication actionnable", level=2)
    document.add_paragraph(
        "Pour chaque proposition, l’utilisateur doit comprendre pourquoi la cellule est retenue, pourquoi les alternatives sont moins bien classées et quelle vérification reste nécessaire. Les détails comme TF-IDF, embeddings et RRF peuvent être proposés dans un volet avancé sans encombrer la vue principale."
    )
    document.add_heading("5.3 Conserver la décision humaine", level=2)
    document.add_paragraph(
        "Le système propose, explique et signale les risques. L’analyste valide, corrige, refuse ou demande une nouvelle recherche. Chaque décision doit conserver l’utilisateur, la date, la valeur précédente, la nouvelle valeur et la justification éventuelle."
    )
    document.add_heading("5.4 Autoriser une abstention sûre", level=2)
    document.add_paragraph(
        "Lorsque les éléments disponibles ne permettent pas une conclusion fiable, l’interface doit afficher clairement qu’aucun candidat suffisamment sûr n’a été identifié. Une abstention explicite est préférable à une réponse artificiellement précise."
    )

    document.add_heading("6. Accompagnement des analystes moins expérimentés", level=1)
    document.add_paragraph(
        "L’interface peut devenir un support d’apprentissage en donnant accès, sans interrompre le travail, à la définition des métriques et des expressions propres au projet."
    )
    bullets(document, (
        "Définition courte puis explication détaillée sur demande.",
        "Unité attendue, plage indicative et formule économique habituelle.",
        "Explication contextuelle des acronymes comme EPC, WHT, DSCR ou gearing.",
        "Distinction visible entre une définition institutionnelle et une explication du LLM.",
    ))

    document.add_heading("7. Gestion des traitements longs", level=1)
    document.add_paragraph(
        "L’extraction doit être exécutée en arrière-plan. L’utilisateur doit pouvoir quitter la page puis revenir sans interrompre le traitement. Une progression fondée sur les étapes réellement terminées doit remplacer une simple animation d’attente."
    )
    callout(document, "Exemple de progression", "Classeur téléversé → catalogue construit → recherche 18/22 → contrôles terminés → validation disponible.")

    document.add_heading("8. Sécurité et confidentialité", level=1)
    document.add_paragraph("L’accès à distance doit intégrer les exigences de sécurité dès le premier prototype utilisable.")
    bullets(document, (
        "Connexion nominative et gestion des rôles.",
        "Chiffrement HTTPS et chiffrement du stockage.",
        "Séparation des projets et contrôle des autorisations.",
        "Journal des consultations, validations et modifications.",
        "Sauvegardes et politique de conservation des fichiers.",
        "Transmission au LLM du contexte strictement nécessaire.",
        "Possibilité de désactiver le LLM pour les dossiers sensibles.",
    ))

    document.add_heading("9. Écrans à concevoir en priorité", level=1)
    priorities = (
        ("1", "Connexion et tableau de bord", "Accéder aux projets, traitements et rapports récents"),
        ("2", "Création du projet", "Déposer le classeur et renseigner le contexte"),
        ("3", "Suivi de l’extraction", "Voir la progression, les alertes et les erreurs"),
        ("4", "Validation des métriques", "Comparer les candidats et corriger valeur et unité"),
        ("5", "Analyse et benchmarks", "Examiner calculs, comparables, écarts et sources"),
        ("6", "Rapports et historique", "Télécharger les livrables et retrouver les versions"),
        ("7", "Administration", "Gérer utilisateurs, référentiels et banque de benchmarks"),
    )
    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    for index, value in enumerate(("Ordre", "Écran", "Fonction principale")):
        shade(table.rows[0].cells[index], NAVY)
        set_cell(table.rows[0].cells[index], value, True, "FFFFFF")
    for row_index, row in enumerate(priorities):
        cells = table.add_row().cells
        for index, value in enumerate(row):
            set_cell(cells[index], value, bold=index in (0, 1))
            if row_index % 2:
                shade(cells[index], GREY)

    document.add_heading("10. Démarche de conception recommandée", level=1)
    document.add_paragraph(
        "Avant de développer toute l’application, il est préférable de réaliser des maquettes des écrans essentiels et de les tester avec des analystes de niveaux différents. Cette démarche permettra de vérifier que les informations facilitent réellement la décision sans créer de surcharge cognitive."
    )
    bullets(document, (
        "Observer le processus actuel et les décisions réelles de l’analyste.",
        "Créer des maquettes du tableau de bord et de la validation.",
        "Les tester avec un analyste expérimenté et un analyste moins expérimenté.",
        "Mesurer le temps de validation, les erreurs et les demandes d’explication.",
        "Développer d’abord le dépôt, l’extraction et la validation.",
        "Ajouter ensuite l’analyse, les benchmarks, les rapports et l’administration.",
    ))

    document.add_heading("11. Décisions à discuter avec le superviseur", level=1)
    bullets(document, (
        "Qui doit accéder à l’application et avec quels rôles ?",
        "L’hébergement sera-t-il interne ou dans un nuage autorisé ?",
        "Les classeurs peuvent-ils être transmis à un service LLM externe ?",
        "Quelle durée de conservation appliquer aux modèles et aux rapports ?",
        "Quelles fonctions sont indispensables dans le premier prototype web ?",
        "Quel niveau de disponibilité et de support est attendu ?",
    ))

    document.add_heading("Conclusion", level=1)
    document.add_paragraph(
        "La future interface ne doit pas être une simple enveloppe graphique autour du script. Elle doit organiser la collaboration entre l’analyste, les règles déterministes, la banque de benchmarks et le LLM. La première priorité de conception est l’écran de validation des métriques, car il concentre la confiance, l’explication, la correction et la traçabilité."
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
