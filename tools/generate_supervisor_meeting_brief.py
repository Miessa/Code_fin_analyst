# -*- coding: utf-8 -*-
"""Génère le support professionnel de la réunion de suivi ARSEL."""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "outputs" / "Point_avancement_superviseur_ARSEL.docx"
BLUE = "17365D"
LIGHT_BLUE = "D9EAF7"
GREEN = "E2F0D9"
AMBER = "FFF2CC"
GREY = "F2F2F2"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text, bold=False, color=None):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Aptos"
    run.font.size = Pt(9)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def configure(document):
    section = document.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(2.1)
    section.right_margin = Cm(2.1)

    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    for name, size, color in (("Title", 25, BLUE), ("Heading 1", 17, BLUE), ("Heading 2", 13, BLUE)):
        style = document.styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("Projet ARSEL  |  Point d’avancement avec le superviseur")
    run.font.name = "Aptos"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(100, 100, 100)


def add_bullets(document, items):
    for item in items:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.add_run(item)


def add_callout(document, title, text, fill=LIGHT_BLUE):
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    shade(cell, fill)
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(title + "\n")
    run.bold = True
    run.font.color.rgb = RGBColor.from_string(BLUE)
    run.font.size = Pt(11)
    paragraph.add_run(text)
    document.add_paragraph()


def add_priority_table(document):
    rows = [
        ("1", "Mesurer la précision", "Jeu de vérité terrain multi-modèles; Top-1, Recall@k, précision de la valeur et de l’unité", "Immédiate"),
        ("2", "Fiabiliser les unités", "Normalisation valeur, devise, échelle, unité, périodicité, année et périmètre", "Immédiate"),
        ("3", "Renforcer les règles métier", "Formules plausibles, exclusions, structures attendues et contrôles de cohérence", "Élevée"),
        ("4", "Améliorer la validation", "Vue simultanée des candidats, contexte Excel, formule, unité et justification du score", "Élevée"),
        ("5", "Sécuriser les benchmarks", "Conversions monétaires, années de référence, périmètres et échantillons comparables", "Élevée"),
        ("6", "Enrichir l’analyse par LLM", "Interprétation après calculs déterministes, avec sources et garde-fous", "Ultérieure"),
        ("7", "Enrichir l’étape 0", "Définition dynamique des acronymes et expressions propres au projet", "Ultérieure"),
    ]
    table = document.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    headers = ("Ordre", "Chantier", "Résultat attendu", "Priorité")
    for index, header in enumerate(headers):
        shade(table.rows[0].cells[index], BLUE)
        set_cell_text(table.rows[0].cells[index], header, bold=True, color="FFFFFF")
    for order, project, result, priority in rows:
        cells = table.add_row().cells
        values = (order, project, result, priority)
        for index, value in enumerate(values):
            set_cell_text(cells[index], value, bold=index in (0, 1))
            if int(order) % 2 == 0:
                shade(cells[index], GREY)


def build():
    document = Document()
    configure(document)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.style = document.styles["Title"]
    title.add_run("Point d’avancement du projet ARSEL")
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Support de réunion avec le superviseur\nÉtat fonctionnel, limites et priorités de développement")
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor.from_string(BLUE)
    document.add_paragraph()
    add_callout(
        document,
        "Message principal",
        "Le prototype couvre désormais la chaîne allant du modèle financier Excel au rapport d’analyse. Le prochain enjeu n’est plus seulement fonctionnel : il consiste à mesurer et augmenter la fiabilité sur plusieurs modèles, en particulier pour les unités et les ambiguïtés métier.",
        GREEN,
    )

    document.add_heading("1. Finalité du projet", level=1)
    document.add_paragraph(
        "Le projet vise à réduire le temps consacré à l’examen des modèles financiers tout en maintenant le contrôle de l’analyste. Le système recherche les principales hypothèses dans un classeur Excel, présente les résultats à la validation, calcule des indicateurs financiers et compare le projet à une banque de références. Il s’agit d’un outil d’aide à l’analyse et non d’un remplacement du jugement professionnel."
    )

    document.add_heading("2. Ce qui fonctionne déjà", level=1)
    document.add_heading("2.1 Extraction hybride des métriques", level=2)
    document.add_paragraph(
        "Le moteur parcourt des classeurs financiers complexes et construit un catalogue de libellés, de valeurs et de contextes. Pour retrouver les cellules pertinentes, il combine la recherche lexicale, TF-IDF et les embeddings. Les classements sont fusionnés avec la méthode RRF, puis enrichis par l’analyse de la structure Excel, des formules et des dépendances entre cellules."
    )
    add_bullets(document, [
        "Collecte de plusieurs milliers de libellés sans chargement répété du classeur.",
        "Présélection combinant correspondance exacte, proximité lexicale et proximité sémantique.",
        "Détection des agrégats et préférence pour les totaux lorsque la métrique l’exige.",
        "Routage conditionnel vers le LLM uniquement pour les cas réellement ambigus.",
        "Conservation d’une shortlist déterministe lorsque le service LLM est indisponible.",
    ])

    document.add_heading("2.2 Règles métier et contrôles de cohérence", level=2)
    document.add_paragraph(
        "Le système ne repose plus uniquement sur la ressemblance textuelle. Des règles métier permettent notamment de distinguer un coût total d’un coût de financement isolé, de préférer une durée annuelle à une observation trimestrielle et d’éviter de confondre une maturité de dette avec une durée de concession. Les valeurs nulles sont signalées comme aberrantes lorsqu’elles ne représentent pas une hypothèse financière crédible."
    )
    add_callout(document, "Exemple de transformation déterministe", "Lorsqu’une cellule indique explicitement une indisponibilité de 10 %, la disponibilité correspondante est calculée selon Disponibilité = 1 − Indisponibilité, soit 90 %.")

    document.add_heading("2.3 Détection et structuration des unités", level=2)
    document.add_paragraph(
        "Une unité peut apparaître à gauche ou à droite de la valeur, dans une ligne supérieure, dans un en-tête ou dans le format de la cellule. Le moteur explore maintenant ces différentes positions et sépare la devise, l’échelle, l’unité physique et la périodicité. Une unité contradictoire ne supprime pas une cellule par ailleurs crédible; elle est conservée comme conflit à vérifier."
    )
    add_bullets(document, [
        "Exemple monétaire : EUR'000s/an devient EUR, échelle 1 000, périodicité annuelle.",
        "Exemple tarifaire : EUR/kW/month devient EUR par kW et par mois.",
        "Exemple de ratio : le DSCR est attendu comme multiple, généralement en x, et non comme pourcentage.",
    ])

    document.add_heading("2.4 Validation par l’analyste", level=2)
    document.add_paragraph(
        "L’analyste peut accepter la proposition, sélectionner une alternative, corriger la valeur, corriger l’unité, modifier les deux simultanément, passer une métrique ou revenir à la précédente. Le résultat est enregistré avec sa source, sa valeur, son unité et la décision prise, ce qui assure la continuité entre l’extraction et l’analyse financière."
    )

    document.add_heading("2.5 Banque de benchmarks et analyse financière", level=2)
    document.add_paragraph(
        "La banque DuckDB regroupe des projets individuels issus de la Banque mondiale et des références sectorielles issues d’IRENA. Les données sont séparées par technologie et gouvernées au moyen d’une zone de staging, de contrôles qualité, de rapports de déduplication, de snapshots adressés par checksum et d’une promotion contrôlée vers la base active."
    )
    add_bullets(document, [
        "Calcul d’indicateurs dérivés à partir des hypothèses validées.",
        "Sélection déterministe de projets comparables selon la technologie, la région, la période et les données disponibles.",
        "Production d’une analyse en JSON, Markdown et Word avec manifeste de traçabilité.",
    ])

    document.add_heading("3. Limites actuelles", level=1)
    document.add_paragraph(
        "Le pipeline est fonctionnel de bout en bout, mais sa précision n’est pas encore suffisamment mesurée sur une diversité de modèles financiers pour permettre une confiance automatique généralisée. Les limites doivent être présentées comme les prochains objets de validation, et non comme des fonctions absentes."
    )
    add_bullets(document, [
        "Certains libellés restent ambigus ou désignent plusieurs périmètres financièrement légitimes.",
        "Les unités peuvent être éloignées, partagées par une zone ou exprimées implicitement.",
        "Les profils temporels nécessitent encore une interprétation plus robuste de leur fréquence et de leur période de référence.",
        "Les règles ont besoin d’être éprouvées sur plusieurs modèles pour confirmer leur généralisation.",
        "Le LLM reste soumis aux délais, quotas et indisponibilités du fournisseur.",
        "Les benchmarks doivent être homogénéisés en devise, année monétaire, unité et périmètre avant comparaison.",
        "La ligne de commande ne permet pas encore une comparaison visuelle optimale de tous les candidats.",
    ])

    document.add_heading("4. Priorités recommandées", level=1)
    add_priority_table(document)

    document.add_heading("5. Comment mesurer le prochain progrès", level=1)
    document.add_paragraph(
        "La prochaine étape structurante consiste à constituer un jeu de vérité terrain sur plusieurs modèles. Pour chaque métrique, ce jeu précisera la cellule attendue, la valeur, l’unité, le périmètre et les alternatives acceptables. La performance devra être mesurée séparément pour la cellule, la valeur et l’unité."
    )
    add_callout(
        document,
        "Indicateurs proposés",
        "Top-1 = nombre de métriques correctes au premier rang / nombre de métriques évaluées.\nRecall@k = nombre de métriques dont la bonne cellule apparaît parmi les k premiers candidats / nombre de métriques évaluées.\nTaux de succès complet = cellule correcte + valeur correcte + unité correcte + périmètre correct.",
        AMBER,
    )

    document.add_heading("6. Démonstration recommandée pendant la réunion", level=1)
    document.add_paragraph(
        "Une démonstration courte permettra de rendre l’avancement plus concret qu’une présentation exclusivement architecturale. Elle peut être organisée autour du scénario suivant :"
    )
    steps = [
        "Lancer l’extraction sur un modèle et montrer trois métriques représentatives.",
        "Afficher la proposition, ses alternatives, son unité et les justifications du score.",
        "Corriger une valeur ou une unité pendant la validation analyste.",
        "Ouvrir le registre des hypothèses validées et montrer la provenance conservée.",
        "Lancer la Phase 2, puis présenter une comparaison provenant de DuckDB.",
        "Ouvrir le rapport Word final et expliquer la séparation entre calcul déterministe et interprétation future par LLM.",
    ]
    for index, step in enumerate(steps, 1):
        document.add_paragraph(f"{index}. {step}")

    document.add_heading("7. Conclusion proposée", level=1)
    document.add_paragraph(
        "À ce stade, nous disposons d’un prototype complet et fonctionnel, allant du classeur Excel au rapport financier. Les améliorations récentes ont renforcé la recherche des métriques, diminué la dépendance au LLM, ajouté une gestion structurée des unités et mis en place une banque de benchmarks gouvernée. La priorité est maintenant de mesurer la précision sur plusieurs modèles, de fiabiliser les unités et de consolider les règles métier. Une fois cette base stabilisée, l’interface utilisateur et l’analyse approfondie par LLM pourront apporter une valeur supplémentaire sans fragiliser les calculs."
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
